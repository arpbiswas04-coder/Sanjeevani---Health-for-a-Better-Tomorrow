from datetime import datetime,timezone,timedelta
from io import BytesIO
import pytest
from openpyxl import load_workbook
from tests.test_platform import seed,receipt
pytestmark=pytest.mark.asyncio


async def test_reports_exports_and_audit(api):
    client,_,_=api
    ids=await seed(client)
    await client.post('/api/v1/inventory/receive',json=receipt(ids))
    for kind in ('stock','expiry','transfers','procurement','staff','beds','emergency'):
        response=await client.get('/api/v1/reports/'+kind)
        assert response.status_code==200,response.text
    csv=await client.get('/api/v1/reports/stock',params={'format':'csv'})
    assert csv.status_code==200 and b'quantity' in csv.content
    xlsx=await client.get('/api/v1/reports/stock',params={'format':'xlsx'})
    sheet=load_workbook(BytesIO(xlsx.content),read_only=True).active
    assert len(list(sheet.rows))==2
    pdf=await client.get('/api/v1/reports/stock',params={'format':'pdf'})
    assert pdf.content.startswith(b'%PDF')
    logs=(await client.get('/api/v1/audit-logs')).json()['data']
    assert 'request_id' in logs[-1]['details']
    assert 'X-Request-ID' in pdf.headers
    assert (await client.get('/api/v1/reports/stock',params={'limit':10000})).status_code==422


async def test_sync_retries_conflicts_and_pagination(api):
    client,_,_=api
    ids=await seed(client)
    now=datetime.now(timezone.utc)
    payload={'idempotency_key':'sync-1','items':[{'kind':'footfall','payload':{'facility_id':ids['facility_id'],'day':str(now.date()),'category':'outpatient','count':4}}]}
    first=await client.post('/api/v1/sync/push',json=payload)
    assert first.status_code==200,first.text
    assert (await client.post('/api/v1/sync/push',json=payload)).json()==first.json()
    assert (await client.post('/api/v1/sync/push',json={**payload,'idempotency_key':'stale'})).status_code==409
    response=await client.get('/api/v1/sync/pull',params={'since':(now-timedelta(days=1)).isoformat()})
    assert response.status_code==200
    assert len(response.json()['data']['items'])==1
    assert response.json()['data']['items'][0]['version']==1


async def test_formula_injection_escaped():
    from app.services.exports import export
    content,_=export([{'name':'=HYPERLINK("bad")'}],'xlsx')
    cell=load_workbook(BytesIO(content)).active['A2']
    assert cell.data_type=='s' and cell.value.startswith("'")
